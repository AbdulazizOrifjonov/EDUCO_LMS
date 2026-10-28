import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Color Palette
NAVY = RGBColor(18, 61, 120)        # #123D78 Brand Navy
NAVY_DARK = RGBColor(10, 35, 71)    # #0A2347 Deep Navy
RED = RGBColor(241, 45, 58)         # #F12D3A Brand Red
WHITE = RGBColor(255, 255, 255)
BG_LIGHT = RGBColor(248, 250, 252)  # #F8FAFC Slate Light
CARD_BG = RGBColor(255, 255, 255)
CARD_BORDER = RGBColor(226, 232, 240)
TEXT_DARK = RGBColor(15, 23, 42)    # #0F172A
TEXT_MUTED = RGBColor(100, 116, 139) # #64748B
GREEN = RGBColor(16, 185, 129)      # #10B981
BLUE = RGBColor(37, 99, 235)        # #2563EB
AMBER = RGBColor(245, 158, 11)      # #F59E0B
PURPLE = RGBColor(139, 92, 246)     # #8B5CF6

def add_header(slide, title_text, category_text="EDUCO LMS QO'LLANMASI"):
    """Adds a standard professional header with category badge"""
    # Category badge
    cat_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.8), Inches(0.35))
    cat_box.fill.solid()
    cat_box.fill.fore_color.rgb = RGBColor(239, 246, 255)
    cat_box.line.color.rgb = BLUE
    cat_box.line.width = Pt(1)
    tf = cat_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = category_text
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = BLUE
    p.alignment = PP_ALIGN.CENTER

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.8))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = NAVY

    # Decorative red line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.55), Inches(1.2), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = RED
    line.line.fill.background()

def add_card(slide, left, top, width, height, title, items, badge_text="", badge_color=BLUE, bg_color=CARD_BG, border_color=CARD_BORDER):
    """Draws a clean modern card with bullet points or descriptions"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)

    # Content
    tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), height - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True

    if badge_text:
        p_badge = tf.paragraphs[0]
        p_badge.text = badge_text.upper()
        p_badge.font.size = Pt(9)
        p_badge.font.bold = True
        p_badge.font.color.rgb = badge_color
        p_badge.space_after = Pt(4)
        p_title = tf.add_paragraph()
    else:
        p_title = tf.paragraphs[0]

    p_title.text = title
    p_title.font.size = Pt(14)
    p_title.font.bold = True
    p_title.font.color.rgb = NAVY
    p_title.space_after = Pt(8)

    for item in items:
        p = tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # ==========================================
    # SLIDE 1: Title Slide (Muqova)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    # Red accent stripe on top
    stripe = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = RED
    stripe.line.fill.background()

    # Brand Title Box
    s1_tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(3.8))
    s1_tf = s1_tb.text_frame
    s1_tf.word_wrap = True

    p0 = s1_tf.paragraphs[0]
    p0.text = "EDUCO.LMS"
    p0.font.size = Pt(28)
    p0.font.bold = True
    p0.font.color.rgb = RED
    p0.space_after = Pt(12)

    p1 = s1_tf.add_paragraph()
    p1.text = "O‘quv Markazini Boshqarish Tizimi"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(10)

    p2 = s1_tf.add_paragraph()
    p2.text = "To‘liq Foydalanish Qo‘llanmasi va Tizim Taqdimoti"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(191, 219, 254)
    p2.space_after = Pt(24)

    p3 = s1_tf.add_paragraph()
    p3.text = "Boshliq (Owner)  |  Administrator  |  O‘qituvchi  |  Talaba"
    p3.font.size = Pt(13)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(226, 232, 240)

    # Footer banner
    footer_box = s1.shapes.add_textbox(Inches(1.2), Inches(6.2), Inches(10.9), Inches(0.8))
    f_tf = footer_box.text_frame
    f_p = f_tf.paragraphs[0]
    f_p.text = "Toshkent — 2026  •  Rasmiy Yo‘riqnoma"
    f_p.font.size = Pt(11)
    f_p.font.color.rgb = RGBColor(148, 163, 184)

    # ==========================================
    # SLIDE 2: Tizim haqida umumiy ma'lumot
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "1. EDUCO LMS Tizimi Haqida Umumiy Ma'lumot")

    add_card(
        s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.2),
        "Yagona Boshqaruv",
        [
            "O‘quv markazining barcha jarayonlarini 100% bitta tizimga jamlaydi.",
            "Talabalar va o‘qituvchilar hisobi qulay jadval ko‘rinishida olib boriladi.",
            "Qog‘ozbozlik va chalkashliklarga to‘liq chek qo‘yadi.",
            "Barcha ma'lumotlar real vaqt rejimida saqlanadi va yangilanadi."
        ],
        badge_text="Avtomatlashtirish",
        badge_color=BLUE
    )

    add_card(
        s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.2),
        "Xavfsiz Rollar Tizimi",
        [
            "4 xil darajadagi foydalanuvchi roli (Boshliq, Admin, Ustoz, Talaba).",
            "Har bir rol faqat o‘ziga tegishli ma'lumotlarni ko‘ra oladi.",
            "Oylik maoshlar va maxfiy audit faqat Boshliq (Owner)ga ko‘rinadi.",
            "Parollarni markazlashgan boshqarish va xavfsiz tiklash."
        ],
        badge_text="Xavfsizlik",
        badge_color=PURPLE
    )

    add_card(
        s2, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.2),
        "Zamonaviy Interfeys",
        [
            "Telefon va kompyuter ekranlariga 100% moslashuvchan (Responsive).",
            "4 ta til qo‘llab-quvvatlanadi: O‘zbek, Rus, Ingliz, Arab.",
            "Kunduzgi va Tungi (Dark Mode) zamonaviy dizayn rejimlari.",
            "Telegram uslubidagi video va audio xabarlar almashinuvi."
        ],
        badge_text="Qulaylik",
        badge_color=GREEN
    )

    # Bottom alert
    bot_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.6), Inches(0.7))
    bot_card.fill.solid()
    bot_card.fill.fore_color.rgb = RGBColor(239, 246, 255)
    bot_card.line.color.rgb = BLUE
    b_tf = bot_card.text_frame
    b_p = b_tf.paragraphs[0]
    b_p.text = "💡 Asosiy maqsad: Markaz ma'muriyati, o‘qituvchilar va o‘quvchilar o‘rtasidagi masofani qisqartirish va sifatni oshirish."
    b_p.font.size = Pt(11)
    b_p.font.bold = True
    b_p.font.color.rgb = NAVY
    b_p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 3: Tizimga Kirish (Login)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "2. Tizimga Kirish (Login) Qoidalari va Format")

    add_card(
        s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
        "Telefon Raqam Bilan Kirish (+998)",
        [
            "Kirish maydonida '+998 ' prefiksi avtomatik ravishda yozilgan bo‘ladi.",
            "Raqamlar yozilishi bilan avtomatik bo‘shliqlar (probel) qo‘yiladi: masalan, +998 90 123 45 67.",
            "Foydalanuvchi xohlagan formatda yozsa ham (probel bilan yoki probelsiz), tizim uni to‘g‘ri taniydi.",
            "Ota-ona telefon raqami orqali ham o‘quvchi kabinetiga kirish imkoni mavjud.",
            "Login qilishda qat'iy tekshiruv va avtomatik xavfsizlik filtri ishlaydi."
        ],
        badge_text="Avtomatik Format",
        badge_color=GREEN
    )

    add_card(
        s3, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8),
        "Login va Parol Tekshiruvi",
        [
            "Username yoki Email orqali ham to‘g‘ridan-to‘g‘ri kirish mumkin (masalan, 'owner', 'admin').",
            "Parol ko‘rinishi uchun ko‘zcha (Eye icon) tugmasidan foydalanish mumkin.",
            "Kirish tugmasi bosilishi bilan tizim foydalanuvchining ROLINI aniqlaydi.",
            "Foydalanuvchi avtomatik ravishda O‘Z DASHBOARDIGA yo‘naltiriladi:",
            "   • Boshliq -> /owner/dashboard",
            "   • Administrator -> /admin/dashboard",
            "   • O‘qituvchi -> /teacher/dashboard",
            "   • Talaba -> /student/dashboard"
        ],
        badge_text="Rolga Yo‘naltirish",
        badge_color=BLUE
    )

    # ==========================================
    # SLIDE 4: 4 Asosiy Rol Matrixi
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "3. Rollar Ierarxiyasi va Ularning Vakolatlari")

    add_card(
        s4, Inches(0.8), Inches(1.8), Inches(2.7), Inches(5.0),
        "Boshliq (Owner)",
        [
            "Tizimning to‘liq egasi.",
            "Faqat 1 ta bo‘ladi.",
            "Admin, O‘qituvchi, Talaba yaratish.",
            "Oylik maoshlar (Payroll).",
            "Xavfsizlik auditi (Audit log).",
            "Barcha moliya va hisobotlar.",
            "Markaz sozlamalarini o‘zgartirish."
        ],
        badge_text="Maksimal Huquq",
        badge_color=RED
    )

    add_card(
        s4, Inches(3.7), Inches(1.8), Inches(2.7), Inches(5.0),
        "Administrator",
        [
            "Markaz boshqaruvchisi.",
            "O‘qituvchi va talaba yaratish.",
            "Guruhlar va kurslar ochish.",
            "Talabalarni guruhga qabul qilish.",
            "To‘lovlarni (Tuition) qayd etish.",
            "Dars jadvallarini tuzish.",
            "Davomat hisobotlarini ko‘rish."
        ],
        badge_text="Boshqaruvchi",
        badge_color=PURPLE
    )

    add_card(
        s4, Inches(6.6), Inches(1.8), Inches(2.7), Inches(5.0),
        "O‘qituvchi",
        [
            "Faqat o‘z guruhlarini ko‘radi.",
            "Har darsda davomat olish.",
            "1 dan 5 gacha baholash.",
            "Mavzu va izoh qoldirish.",
            "Uyga vazifa yaratish va tekshirish.",
            "O‘quv materiallari yuklash.",
            "Guruh chati orqali muloqot."
        ],
        badge_text="Ta'lim Jarayoni",
        badge_color=BLUE
    )

    add_card(
        s4, Inches(9.5), Inches(1.8), Inches(2.7), Inches(5.0),
        "Talaba / O‘quvchi",
        [
            "O‘zining dars jadvalini ko‘rish.",
            "Olingan baholar tarixi.",
            "Davomat statistikasi (foizda).",
            "Uyga vazifalarni yuklab olish.",
            "Yechilgan vazifani topshirish.",
            "Guruh chatida qatnashish.",
            "To‘lov holatini kuzatish."
        ],
        badge_text="O‘quvchi",
        badge_color=GREEN
    )

    # ==========================================
    # SLIDE 5: Yangi Rol va Foydalanuvchi Ochish Tartibi
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "4. Yangi Akkaunt va Rol Yaratish Qadam-baqadam")

    add_card(
        s5, Inches(0.8), Inches(1.8), Inches(6.8), Inches(5.0),
        "Akkaunt Yaratishning 5 Ta Aniq Qadami",
        [
            "1-QADAM: Boshliq yoki Admin panelida chap menyudan 'Foydalanuvchilar' (Users) bo‘limiga o‘ting.",
            "2-QADAM: Sahifaning yuqori o‘ng burchagidagi ko‘k rangli '+ Foydalanuvchi yaratish' tugmasini bosing.",
            "3-QADAM: Oynada Ism, Familiya va Telefon raqamini kiriting (+998 prefiksi tayyor turadi).",
            "4-QADAM: Rolni tanlang — Administrator (faqat Owner ocha oladi), O‘qituvchi yoki Talaba.",
            "5-QADAM: Kamida 6 ta belgidan iborat Parol kiriting (masalan, 'bobur2026' yoki 'educo123').",
            "NATIJA: 'Saqlash' tugmasi bosilishi bilan akkaunt yaratiladi va o‘sha zahotiyoq login qila oladi!"
        ],
        badge_text="Amaliy Qo‘llanma",
        badge_color=BLUE
    )

    add_card(
        s5, Inches(7.9), Inches(1.8), Inches(4.6), Inches(5.0),
        "Muhim Qoidalar & Tavsiyalar",
        [
            "⚡ Parol talabi: Kamida 6 ta belgi bo‘lishi kifoya. Murakkab belgilar majburiy emas.",
            "⚡ Telefon raqam: Har bir foydalanuvchining telefon raqami login vazifasini ham bajaradi.",
            "⚡ Talaba ochilayotganda: Otasi yoki onasining telefon raqami kiritilsa, ota-ona ham tizimga kira oladi.",
            "⚡ Guruhga biriktirish: Talabani yaratish paytidayoq mavjud faol guruhga avtomatik qo‘shish mumkin.",
            "⚡ Boshliq (Owner): Tizimda faqat bitta bo‘ladi, uni qayta yaratib bo‘lmaydi."
        ],
        badge_text="Eslatma",
        badge_color=AMBER,
        bg_color=RGBColor(255, 251, 235),
        border_color=RGBColor(253, 230, 138)
    )

    # ==========================================
    # SLIDE 6: Parollar Sahifasi (Passwords Management)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "5. 'Parollar' Bo‘limi va Ularni Boshqarish")

    add_card(
        s6, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
        "Parollarni Ko‘rish",
        [
            "Boshliq va Admin menyusidagi 'Parollar' sahifasida barcha foydalanuvchilar ro‘yxati turadi.",
            "Har bir xodim va o‘quvchining amaldagi paroli xavfsiz holda ko‘rsatiladi.",
            "Ko‘zcha tugmasini bosib barcha parollarni bittada ochish yoki yashirish mumkin."
        ],
        badge_text="Nazorat",
        badge_color=BLUE
    )

    add_card(
        s6, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8),
        "Bir Bosishda Nusxalash",
        [
            "Har bir qatorda 'Nusxa olish' (Copy) tugmasi mavjud.",
            "Tugmani bosishingiz bilan login va parol xotiraga olinadi.",
            "O‘qituvchi yoki o‘quvchiga Telegram orqali bitta tugma bilan tayyor ma'lumotni jo‘natish juda qulay."
        ],
        badge_text="Tezkorlik",
        badge_color=GREEN
    )

    add_card(
        s6, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8),
        "Yangi Parol O‘rnatish",
        [
            "Agar o‘quvchi yoki xodim parolini unutib qo‘ysa, 'Tahrirlash' tugmasini bosing.",
            "Yangi parolni yozing (kamida 6 ta belgi) va 'Saqlash' tugmasini bosing.",
            "Parol darhol yangilanadi va foydalanuvchi yangi parol bilan kabinetga kiradi."
        ],
        badge_text="Parol Tiklash",
        badge_color=PURPLE
    )

    # ==========================================
    # SLIDE 7: Kurslar va Guruhlar Boshqaruvi
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "6. Kurslar va Guruhlarni Shakllantirish")

    add_card(
        s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Kurslar (Courses)",
        [
            "Markazdagi barcha ta'lim yo‘nalishlari: Frontend, Python, IELTS, Arab tili, Grafik dizayn va h.k.",
            "Har bir kurs uchun oylik to‘lov narxi (Monthly fee) va yo‘nalish kodi belgilanadi.",
            "Kursga tegishli faol guruhlar soni va umumiy o‘quvchilar soni avtomatik hisoblab boriladi.",
            "Eski yoki tugatilgan kurslarni arxivlash imkoniyati."
        ],
        badge_text="Ta'lim Dasturlari",
        badge_color=BLUE
    )

    add_card(
        s7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Guruhlar (Groups)",
        [
            "Guruh nomi, kodi, biriktirilgan kurs va o‘qituvchi tanlanadi.",
            "Dars xonasi (Room) va talabalar maksimal sig‘imi (Capacity) belgilanadi.",
            "Talabalarni guruhga qo‘shish (Enrollment) birgina tugma orqali bajariladi.",
            "Talabani boshqa guruhga ko‘chirish (Transfer) tarixi to‘liq saqlanib qolinadi.",
            "Guruhning boshlanish va tugash sanalari bo‘yicha hisobot."
        ],
        badge_text="Sinflar Boshqaruvi",
        badge_color=GREEN
    )

    # ==========================================
    # SLIDE 8: Dars Jadvali va Davomat
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "7. Interaktiv Dars Jadvali va Kunlik Davomat")

    add_card(
        s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Haftalik Dars Jadvali",
        [
            "Dushanbadan Yakshanbagacha barcha darslar ko‘rgazmali jadvalda ko‘rinadi.",
            "Har bir guruh o‘zining maxsus rangida ajralib turadi.",
            "Xona to‘qnashuvini aniqlash: Agar bir xonaga bir vaqtda 2 ta dars qo‘yilsa, tizim ogohlantiradi.",
            "O‘qituvchi to‘qnashuvini aniqlash: O‘qituvchining bir vaqtda ikki joyda darsi bo‘lishi oldi olinadi."
        ],
        badge_text="Timetable",
        badge_color=PURPLE
    )

    add_card(
        s8, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Tezkor Davomat Tizimi",
        [
            "O‘qituvchi guruhni tanlaydi va o‘quvchilar ro‘yxati chiqadi.",
            "4 xil status bir marta bosishda belgilanadi:",
            "   🟢 Hozir (Present) — darsda qatnashmoqda",
            "   🟡 Kech qoldi (Late) — darsga kechikib keldi",
            "   🔵 Sababli (Excused) — uzrli sabab bilan yo‘q",
            "   🔴 Sababsiz (Unexcused) — sababsiz dars qoldirdi",
            "Sababli qoldirilgan darslar uchun izoh yozish maydoni mavjud.",
            "Davomat oylik maosh va sertifikat berish foizlariga ta'sir qiladi."
        ],
        badge_text="Attendance",
        badge_color=BLUE
    )

    # ==========================================
    # SLIDE 9: Baholar va Uyga Vazifalar
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "8. 5-Ballik Baholash va Uyga Vazifalar")

    add_card(
        s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "5 Ballik Standart Baholash",
        [
            "Har bir talabaga mavzu bo‘yicha 1 dan 5 gacha baho qo‘yiladi:",
            "   5 — A'lo (Yashil belgi)",
            "   4 — Yaxshi (Ko‘k belgi)",
            "   3 — Qoniqarli (Sariq belgi)",
            "   2 — Qoniqarsiz (To‘q sariq belgi)",
            "   1 — Yomon (Qizil belgi)",
            "O‘qituvchi har bir baho uchun o‘quvchiga yo‘naltiruvchi fikr (Feedback) yozishi mumkin.",
            "Talaba o‘z kabinetida barcha baholari dinamikasini ko‘rib boradi."
        ],
        badge_text="Baholash Tizimi",
        badge_color=GREEN
    )

    add_card(
        s9, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Uyga Vazifalar (Homework)",
        [
            "O‘qituvchi vazifa mavzusi, batafsil tavsif va topshirish muddatini (Deadline) belgilaydi.",
            "Vazifaga qo‘shimcha kitob, rasm yoki hujjat fayllarini biriktirish mumkin.",
            "O‘quvchilar o‘z kabinetidan yechilgan vazifani yuborishadi.",
            "O‘qituvchi topshirilgan ishlarni tekshiradi, qabul qiladi yoki qayta topshirishga qaytaradi.",
            "Kechikkan va topshirilmagan vazifalar statistikasi shakllanadi."
        ],
        badge_text="Vazifalar Nazorati",
        badge_color=BLUE
    )

    # ==========================================
    # SLIDE 10: Guruh Chati va Video/Audio Xabarlar
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "9. Guruh Chati, Video va Audio Xabarlar")

    add_card(
        s10, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Guruh Muloqoti",
        [
            "Telegram uslubidagi qulay guruh chati.",
            "Har bir guruh uchun alohida umumiy dars xonasi muloqoti.",
            "Xodimlar uchun (Owner, Admin, Ustozlar) maxsus ichki aloqa chati.",
            "Xabarga javob qaytarish (Reply) va o‘chirish imkoniyati."
        ],
        badge_text="Interaktiv Chat",
        badge_color=BLUE
    )

    add_card(
        s10, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "To‘rtburchak Video Xabar",
        [
            "Maksimal 60 soniyali qisqa video xabar yozish.",
            "To‘g‘ri to‘rtburchak (Rectangular) zamonaviy format.",
            "Kamerani oldi va orqaga o‘girish (Front / Back camera switch) tugmasi.",
            "Videoni oldindan ko‘rish va qayta yozish imkoniyati.",
            "Telefonda ham, kompyuterda ham qotmasdan ishlaydi."
        ],
        badge_text="Video Recorder",
        badge_color=RED
    )

    add_card(
        s10, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Ovozli Xabar va Fayllar",
        [
            "Ovozli xabar (Voice message) yozish va bir zumda yuborish.",
            "Rasm, PDF va o‘quv qo‘llanmalarini yuklash.",
            "Emojilar klaviaturasi orqali jonli muloqot.",
            "Yangi xabarlar haqida real vaqtda ovozli va vizual bildirishnomalar."
        ],
        badge_text="Multimedia",
        badge_color=GREEN
    )

    # ==========================================
    # SLIDE 11: Moliya, Oyliklar va To'lovlar
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "10. Moliya: Oylik Maoshlar (Payroll) va To‘lovlar (Tuition)")

    add_card(
        s11, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Oylik Maoshlar (Faqat Boshliq uchun)",
        [
            "Xodimlar maoshi 3 xil usulda avtomatik hisoblanadi:",
            "   1. Fiks oylik (Fixed Monthly) — har oy belgilangan oylik",
            "   2. Soatbay (Hourly) — ishlangan soatlarga qarab",
            "   3. Darsbay (Per Lesson) — o‘tilgan darslar soniga qarab",
            "Oylik davrlarini ochish (masalan, Oktyabr 2026) va yakunlash.",
            "Bonus va jarimalarni hisobga olgan holda aniq to‘lov ro‘yxati.",
            "Barcha to‘lovlarni Excel (CSV) formatida yuklab olish."
        ],
        badge_text="Xodimlar Maoshi",
        badge_color=PURPLE
    )

    add_card(
        s11, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Talabalar To‘lovlari (Tuition)",
        [
            "O‘quvchilarning to‘lov holati doimiy nazoratda:",
            "   🟢 To‘langan (Paid)",
            "   🟡 Qisman to‘langan (Partial)",
            "   🔴 Qarzdor (Unpaid)",
            "Qabul qilingan to‘lovlarni kassa orqali qayd etish (Naqd, Karta, O‘tkazma).",
            "Qarzdor talabalar bo‘yicha tezkor ogohlantirishlar.",
            "Oylik tushumlar va umumiy aylanma dinamikasi grafiklarda aks etadi."
        ],
        badge_text="O‘quv To‘lovlari",
        badge_color=GREEN
    )

    # ==========================================
    # SLIDE 12: Xulosa va Qisqa Maslahatlar
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    bg12 = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg12.fill.solid()
    bg12.fill.fore_color.rgb = NAVY_DARK
    bg12.line.fill.background()

    # Red accent stripe on top
    stripe12 = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    stripe12.fill.solid()
    stripe12.fill.fore_color.rgb = RED
    stripe12.line.fill.background()

    # Header in slide 12
    t12_box = s12.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.0))
    t12_tf = t12_box.text_frame
    p_t12 = t12_tf.paragraphs[0]
    p_t12.text = "EDUCO LMS Bilan Samarali Ishlash Bo‘yicha 5 Qoida"
    p_t12.font.size = Pt(26)
    p_t12.font.bold = True
    p_t12.font.color.rgb = WHITE

    # 5 Rules Cards Grid
    rules = [
        ("1. Avval Kurs va O‘qituvchini Ochish", "Yangi guruh yaratishdan oldin kurs yo‘nalishini va unga javobgar o‘qituvchini tizimga kiritib oling."),
        ("2. Har Darsda Davomatni Kiritish", "Davomatni dars kuni kechiktirmasdan kiriting. Bu maosh va qarzdorliklar hisobotiga to‘g‘ridan-to‘g‘ri ta'sir qiladi."),
        ("3. Parollarni Xavfsiz Yetkazish", "'Parollar' bo‘limidan foydalanib, yangi ochilgan o‘quvchilarga login va parolni qulay usulda yetkazing."),
        ("4. Telefonda Qulay Foydalanish", "Mobil brauzerda tizim to‘liq moslashgan, gorizontal skroll bo‘lmasligi uchun barcha elementlar tartibga solingan."),
        ("5. Oylik Yakunida Hisobotlarni Yopish", "Har oyning so‘nggi kunida davomat va oylik davrlarini tekshirib, yakuniy hisobotlarni tasdiqlang.")
    ]

    top_pos = Inches(2.0)
    for title, desc in rules:
        r_box = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), top_pos, Inches(11.333), Inches(0.85))
        r_box.fill.solid()
        r_box.fill.fore_color.rgb = RGBColor(18, 61, 120)
        r_box.line.color.rgb = RGBColor(59, 130, 246)
        r_box.line.width = Pt(1)

        r_tf = r_box.text_frame
        r_tf.word_wrap = True
        rp0 = r_tf.paragraphs[0]
        rp0.text = title
        rp0.font.size = Pt(13)
        rp0.font.bold = True
        rp0.font.color.rgb = RGBColor(254, 202, 202)
        rp0.space_after = Pt(2)

        rp1 = r_tf.add_paragraph()
        rp1.text = desc
        rp1.font.size = Pt(10.5)
        rp1.font.color.rgb = RGBColor(241, 245, 249)

        top_pos += Inches(0.95)

    # Save PPTX
    out_dir = r"C:\Users\Abdul\Documents\EDUCO_LMS"
    pptx_path = os.path.join(out_dir, "EDUCO_LMS_Qollanma_Taqdimot.pptx")
    pdf_path = os.path.join(out_dir, "EDUCO_LMS_Qollanma_Taqdimot.pdf")

    prs.save(pptx_path)
    print(f"PPTX muvaffaqiyatli saqlandi: {pptx_path}")

    # Convert to PDF using PowerPoint Application
    try:
        import win32com.client
        ppt_app = win32com.client.Dispatch("PowerPoint.Application")
        presentation = ppt_app.Presentations.Open(pptx_path, WithWindow=False)
        presentation.SaveAs(pdf_path, 32) # 32 = ppSaveAsPDF
        presentation.Close()
        ppt_app.Quit()
        print(f"PDF muvaffaqiyatli yaratildi: {pdf_path}")
    except Exception as e:
        print(f"PDF ga aylantirishda xatolik: {e}")

if __name__ == "__main__":
    build_presentation()
