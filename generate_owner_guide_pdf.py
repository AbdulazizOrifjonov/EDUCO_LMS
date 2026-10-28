import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Color Palette (Dark Mode Brand Aesthetic matching the screenshots)
NAVY_BG = RGBColor(15, 23, 42)       # #0F172A Dark Slate
NAVY_SURFACE = RGBColor(30, 41, 59)  # #1E293B Surface Card
NAVY_BRAND = RGBColor(18, 61, 120)   # #123D78 Brand Navy
RED_BRAND = RGBColor(241, 45, 58)    # #F12D3A Brand Red
BLUE_ACCENT = RGBColor(59, 130, 246) # #3B82F6
CYAN_ACCENT = RGBColor(6, 182, 212)  # #06B6D4
GREEN_ACCENT = RGBColor(16, 185, 129)# #10B981
AMBER_ACCENT = RGBColor(245, 158, 11)# #F59E0B
WHITE = RGBColor(255, 255, 255)
TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8
BORDER_COLOR = RGBColor(51, 65, 85)  # #334155

SCREENSHOTS_DIR = os.path.abspath("./screenshots_owner")

SLIDES_DATA = [
    {
        "img": "01_login.png",
        "num": "01",
        "title": "Tizimga Kirish Sahifasi (Login)",
        "url": "http://localhost:5173/login",
        "badge": "XAVFSIZLIK VA AVTORIZATSIYA",
        "badge_color": BLUE_ACCENT,
        "features": [
            "Telefon raqam (+998), Username yoki Email orqali tezkor kirish.",
            "Kirish maydonida '+998 ' prefiksi avtomatik chiqadi va raqamlar terilishi bilan probellar bilan shakllanadi (+998 90 123 45 67).",
            "Parol ko'rinishini boshqarish uchun ko'zcha (Eye) tugmasi.",
            "Boshliq ma'lumotlari kiritilishi bilan avtomatik Boshliq boshqaruv paneliga (/owner/dashboard) yo'naltiriladi."
        ]
    },
    {
        "img": "02_dashboard.png",
        "num": "02",
        "title": "Boshliq Boshqaruv Paneli (Dashboard)",
        "url": "http://localhost:5173/owner",
        "badge": "ASOSIY PORTAL & MONITORING",
        "badge_color": RED_BRAND,
        "features": [
            "O'quv markazining real vaqtdagi 6 ta asosiy statistik ko'rsatkichi: Jami talabalar, Ustozlar, Administratorlar, Guruhlar, Faol guruhlar va Kurslar.",
            "Bugungi o'tiladigan darslar ro'yxati va dars xonalari taqsimoti.",
            "Haftalik davomat ko'rsatkichlari dinamikasi va umumiy davomat darajasi.",
            "Tezkor amallar tugmalari: Yangi o'quvchi qo'shish, O'qituvchi qo'shish, Guruh yaratish va Hisobotlar."
        ]
    },
    {
        "img": "03_users.png",
        "num": "03",
        "title": "Foydalanuvchilar Boshqaruvi (Users)",
        "url": "http://localhost:5173/owner/users",
        "badge": "XODIMLAR VA O'QUVCHILAR RO'YXATI",
        "badge_color": BLUE_ACCENT,
        "features": [
            "Markazning barcha foydalanuvchilari (Adminlar, O'qituvchilar, Talabalar) to'liq ro'yxati.",
            "Rollar bo'yicha saralash (filter) va Ism/Telefon bo'yicha tezkor qidiruv.",
            "Yangi foydalanuvchi ochish (+ Foydalanuvchi yaratish) tugmasi.",
            "Har bir akkauntni tahrirlash (Edit), parolini yangilash (Key) va faollik holatini (Status) bir bosishda boshqarish."
        ]
    },
    {
        "img": "04_passwords.png",
        "num": "04",
        "title": "Parollar va Xavfsizlik Bo'limi (Passwords)",
        "url": "http://localhost:5173/owner/passwords",
        "badge": "MAXSUS PAROLLAR NAZORATI",
        "badge_color": AMBER_ACCENT,
        "features": [
            "Barcha xodimlar va o'quvchilar parollarini markazlashgan holda ko'rish.",
            "Nusxa olish (Copy) tugmasi: Bitta bosishda login va parolni nusxalab, Telegram orqali xodimga yuborish imkoni.",
            "Avto-parol generatori va yangi parol biriktirish (Tahrirlash tugmasi orqali).",
            "Barcha parollarni bittada ochish yoki yashirish imkoniyati."
        ]
    },
    {
        "img": "05_courses.png",
        "num": "05",
        "title": "Kurslar Boshqaruvi (Courses)",
        "url": "http://localhost:5173/owner/courses",
        "badge": "TA'LIM YO'NALISHLARI",
        "badge_color": GREEN_ACCENT,
        "features": [
            "Barcha o'quv yo'nalishlari: Frontend React, Python Backend, IELTS Intensive, UI/UX Design va boshqalar.",
            "Har bir kursning oylik to'lov summasi (Monthly fee), kodi va yo'nalishi.",
            "Kursga tegishli faol guruhlar soni va umumiy talabalar statistikasi.",
            "Yangi kurs qo'shish va eskirgan kurslarni arxivlash."
        ]
    },
    {
        "img": "06_groups.png",
        "num": "06",
        "title": "Guruhlar Boshqaruvi (Groups)",
        "url": "http://localhost:5173/owner/groups",
        "badge": "SINFLAR VA GURUHLAR",
        "badge_color": CYAN_ACCENT,
        "features": [
            "Guruh nomi, kursi, biriktirilgan o'qituvchi, dars xonasi va o'quvchilar sig'imi (masalan: 16 kishi).",
            "Talabalarni guruhga qabul qilish (Enroll) va boshqa guruhga ko'chirish (Transfer).",
            "Guruhning boshlanish va tugash sanalari, dars kunlari va soatlari.",
            "Faol va arxivlangan guruhlar bo'yicha qulay filtrlash."
        ]
    },
    {
        "img": "07_schedule.png",
        "num": "07",
        "title": "Haftalik Dars Jadvali (Schedule)",
        "url": "http://localhost:5173/owner/schedule",
        "badge": "INTERAKTIV TIMETABLE",
        "badge_color": BLUE_ACCENT,
        "features": [
            "Dushanbadan Yakshanbagacha haftalik darslar jadvali interaktiv vizual ko'rinishda.",
            "Har bir guruh o'zining alohida rangida ajratib ko'rsatiladi.",
            "Xona to'qnashuvi nazorati: Bir xonaga bir vaqtda 2 ta dars belgilanishining oldi olinadi.",
            "O'qituvchi bandligi nazorati: O'qituvchining dars vaqtlari to'g'ri taqsimlanadi."
        ]
    },
    {
        "img": "08_attendance.png",
        "num": "08",
        "title": "Kunlik Davomat Tizimi (Attendance)",
        "url": "http://localhost:5173/owner/attendance",
        "badge": "DAVOMAT VA QATNASHISH",
        "badge_color": GREEN_ACCENT,
        "features": [
            "Kunlik darslar bo'yicha talabalar davomatini tezkor belgilash.",
            "4 xil status: Hozir (Present / Yashil), Kech qoldi (Late / Sariq), Sababli (Excused / Ko'k), Sababsiz (Unexcused / Qizil).",
            "Sababli qoldirilgan darslar uchun izoh yozish imkoniyati.",
            "Oylik davomat hisoboti va umumiy davomat foizlari shakllantirilishi."
        ]
    },
    {
        "img": "09_grades.png",
        "num": "09",
        "title": "5-Ballik Baholash Tizimi (Grades)",
        "url": "http://localhost:5173/owner/grades",
        "badge": "O'ZLASHTIRISH VA BAHOLAR",
        "badge_color": AMBER_ACCENT,
        "features": [
            "Talabalar bilimini 1 dan 5 gacha standart shkalada baholash:",
            "5 (A'lo), 4 (Yaxshi), 3 (Qoniqarli), 2 (Qoniqarsiz), 1 (Yomon).",
            "Har bir baho uchun o'qituvchi tomonidan yo'naltiruvchi fikr (Feedback) kiritish.",
            "O'quvchining davrlar bo'yicha o'rtacha balli va baholar tarixi."
        ]
    },
    {
        "img": "10_homework.png",
        "num": "10",
        "title": "Uyga Vazifalar Nazorati (Homework)",
        "url": "http://localhost:5173/owner/homework",
        "badge": "VAZIFALAR VA TOPSHIRIQLAR",
        "badge_color": BLUE_ACCENT,
        "features": [
            "Yangi uyga vazifa yaratish, mavzu, topshirish muddati (Deadline) belgilash.",
            "Vazifaga kitob, fayl, rasm va qo'shimcha topshiriqlarni biriktirish.",
            "Talabalar tomonidan yuborilgan javoblarni ko'rish, qabul qilish yoki qaytarish.",
            "Topshirilgan, topshirilmagan va kechiktirilgan vazifalar statistikasi."
        ]
    },
    {
        "img": "11_materials.png",
        "num": "11",
        "title": "O'quv Materiallari Markazi (Materials)",
        "url": "http://localhost:5173/owner/materials",
        "badge": "ELEKTRON KUTUBXONA",
        "badge_color": CYAN_ACCENT,
        "features": [
            "Darsliklar, PDF qo'llanmalar, taqdimotlar va dastur kodlarini guruhlarga yuklash.",
            "Guruhlar va kurslar bo'yicha saralangan elektron kutubxona.",
            "Fayl hajmi va yuklangan sanasi bilan qulay kartochkalar.",
            "Talabalar o'z kabinetidan materiallarni bir bosishda yuklab olishi mumkin."
        ]
    },
    {
        "img": "12_announcements.png",
        "num": "12",
        "title": "Markaz E'lonlari (Announcements)",
        "url": "http://localhost:5173/owner/announcements",
        "badge": "XABARNOMALAR VA E'LONLAR",
        "badge_color": RED_BRAND,
        "features": [
            "Markaz bo'yicha umumiy yoki alohida guruhlarga e'lonlar yaratish.",
            "Muhim sanalar, bayramlar, imtihonlar va o'zgarishlar haqida tezkor ogohlantirish.",
            "E'lon chop etilishi bilan barcha foydalanuvchilar ekranida bildirishnoma chiqadi.",
            "Eski e'lonlarni tahrirlash va arxivlash imkoni."
        ]
    },
    {
        "img": "13_chat.png",
        "num": "13",
        "title": "Guruh Chati va Video/Audio Xabarlar",
        "url": "http://localhost:5173/owner/chat",
        "badge": "TELEGRAM USLUBIDAGI ALOQA",
        "badge_color": BLUE_ACCENT,
        "features": [
            "Har bir guruh uchun alohida umumiy dars chati va xodimlar ichki chati.",
            "To'rtburchak Video Xabar (Video Recorder): 60 soniyali video, kamerani old/orqa o'girish tugmasi.",
            "Ovozli xabarlar (Audio Recorder) yozish va bir zumda yuborish.",
            "Rasm, fayl, hujjatlar yuborish, reply va xabarni o'chirish imkoniyati."
        ]
    },
    {
        "img": "14_payroll.png",
        "num": "14",
        "title": "Ish Haqi va Oyliklar (Payroll)",
        "url": "http://localhost:5173/owner/payroll",
        "badge": "FAQAT BOSHLIQ (OWNER) UCHUN",
        "badge_color": RED_BRAND,
        "features": [
            "3 xil oylik hisoblash usuli: Fiks (Fixed Monthly), Soatbay (Hourly), Darsbay (Per Lesson).",
            "Oylik davrlarini ochish (masalan: Oktyabr 2026), hisoblash va yopish.",
            "O'qituvchilarning o'tgan darslari va soatlari bo'yicha avtomatik oylik hisoboti.",
            "Bonuslar va ushlanmalarni kiritish, barcha to'lovlarni Excel (CSV)ga eksport qilish."
        ]
    },
    {
        "img": "15_tuition.png",
        "num": "15",
        "title": "O'qish To'lovlari Nazorati (Tuition)",
        "url": "http://localhost:5173/owner/tuition",
        "badge": "MOLIYA VA KASSA HISOBI",
        "badge_color": GREEN_ACCENT,
        "features": [
            "Talabalarning o'qish to'lovlari holati: To'langan (Paid), Qisman (Partial), Qarzdor (Unpaid).",
            "Kassa orqali yangi to'lovni qabul qilish (Naqd, Karta, Bank o'tkazmasi).",
            "Talabalar bo'yicha to'lovlar tarixi va kvitansiyalar hisobi.",
            "Qarzdor talabalar bo'yicha tezkor eslatmalar va umumiy moliyaviy tushumlar balansi."
        ]
    },
    {
        "img": "16_reports.png",
        "num": "16",
        "title": "Hisobotlar Markazi (Reports)",
        "url": "http://localhost:5173/owner/reports",
        "badge": "KENG QAMROVLI ANALITIKA",
        "badge_color": BLUE_ACCENT,
        "features": [
            "Davomat hisoboti: Qatnashish darajasi va guruhlar bo'yicha qoldirilgan darslar analitikasi.",
            "Baholar hisoboti: Talabalar va kurslar bo'yicha o'rtacha o'zlashtirish ko'rsatkichlari.",
            "O'qituvchilar hisoboti: O'tilgan darslar soati va ustozlar yuklamasi.",
            "Sanalarni tanlash orqali barcha hisobotlarni Excel (CSV) fayl ko'rinishida yuklab olish."
        ]
    },
    {
        "img": "17_audit.png",
        "num": "17",
        "title": "Xavfsizlik Audit Jurnali (Audit Log)",
        "url": "http://localhost:5173/owner/audit",
        "badge": "XAVFSIZLIK VA HARAKATLAR TARIXI",
        "badge_color": AMBER_ACCENT,
        "features": [
            "Faqat Boshliq uchun ochiq: Tizimda sodir bo'lgan har bir amal to'liq qayd etiladi.",
            "Foydalanuvchi yaratilishi, parol o'zgartirilishi, guruhdan o'chirish yoki baho qo'yish aniq vaqti bilan saqlanadi.",
            "Amalni kim bajarganligi va qaysi obyektga ta'sir qilganligi ko'rsatiladi.",
            "Markaz xavfsizligi va shaffofligini ta'minlovchi asosiy nazorat vositasi."
        ]
    },
    {
        "img": "18_notifications.png",
        "num": "18",
        "title": "Bildirishnomalar Markazi (Notifications)",
        "url": "http://localhost:5173/owner/notifications",
        "badge": "REAL VAQTDAGI SIGNALAR",
        "badge_color": CYAN_ACCENT,
        "features": [
            "Markazdagi barcha muhim hodisalar haqidagi real vaqtdagi bildirishnomalar.",
            "Yangi to'lovlar tushishi, dars jadvali o'zgarishi, yangi baholar va guruhga qabullar.",
            "O'qilmagan bildirishnomalar soni va barchasini o'qilgan deb belgilash tugmasi.",
            "Bildirishnomani bosish orqali to'g'ridan-to'g'ri tegishli bo'limga o'tish."
        ]
    },
    {
        "img": "19_settings.png",
        "num": "19",
        "title": "Markaz va Tizim Sozlamalari (Settings)",
        "url": "http://localhost:5173/owner/settings",
        "badge": "KONFIGURATSIYA VA MOSLASHTIRISH",
        "badge_color": BLUE_ACCENT,
        "features": [
            "O'quv markazi nomi, logotipi va vaqt mintaqasi (Asia/Tashkent).",
            "Davomat qoidalari: Avto-davomat vaqti va sukut bo'yicha sabablar.",
            "Til sozlamalari (O'zbek, Rus, Ingliz, Arab tillari o'rtasida bir zumda almashtirish).",
            "Tungi (Dark) va Kunduzgi (Light) rejimlar, profil parolini yangilash."
        ]
    }
]

def build_pdf_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # ==========================================================
    # SLIDE 1: Title Slide
    # ==========================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY_BG
    bg1.line.fill.background()

    # Red Accent Line on Top
    top_line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = RED_BRAND
    top_line.line.fill.background()

    # Title Card Box
    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(10.9), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "EDUCO.LMS  •  BOSHLIQ (OWNER) NIZOMI"
    p0.font.size = Pt(20)
    p0.font.bold = True
    p0.font.color.rgb = RED_BRAND
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "Boshliq Rolining Barcha Sahifalari"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "To'liq Skrinshotlar va Qadam-baqadam Foydalanish Qo'llanmasi"
    p2.font.size = Pt(22)
    p2.font.color.rgb = CYAN_ACCENT
    p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    p3.text = "Boshqaruv Paneli  |  Foydalanuvchilar  |  Parollar  |  Kurslar  |  Guruhlar  |  Davomat  |  Chat  |  Payroll"
    p3.font.size = Pt(13)
    p3.font.color.rgb = TEXT_MUTED

    # Bottom badge
    bb = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(6.2), Inches(10.9), Inches(0.6))
    bb.fill.solid()
    bb.fill.fore_color.rgb = NAVY_SURFACE
    bb.line.color.rgb = BORDER_COLOR
    bbtf = bb.text_frame
    bbp = bbtf.paragraphs[0]
    bbp.text = "Litsenziya: Rasmiy Yo'riqnoma  •  Alisher Qodirov (Boshliq)  •  Toshkent 2026"
    bbp.font.size = Pt(11)
    bbp.font.bold = True
    bbp.font.color.rgb = WHITE
    bbp.alignment = PP_ALIGN.CENTER

    # ==========================================================
    # SLIDES 2-20: Each Owner Page with Crisp Screenshot
    # ==========================================================
    for item in SLIDES_DATA:
        s = prs.slides.add_slide(blank_layout)

        # Slide background
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = NAVY_BG
        bg.line.fill.background()

        # Top Red Stripe
        st = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.08))
        st.fill.solid()
        st.fill.fore_color.rgb = RED_BRAND
        st.line.fill.background()

        # Top Badge (Category)
        badge_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.28), Inches(3.4), Inches(0.32))
        badge_box.fill.solid()
        badge_box.fill.fore_color.rgb = NAVY_SURFACE
        badge_box.line.color.rgb = item["badge_color"]
        badge_box.line.width = Pt(1)
        btf = badge_box.text_frame
        bp = btf.paragraphs[0]
        bp.text = item["badge"]
        bp.font.size = Pt(9.5)
        bp.font.bold = True
        bp.font.color.rgb = item["badge_color"]
        bp.alignment = PP_ALIGN.CENTER

        # URL Badge on right
        url_box = s.shapes.add_textbox(Inches(4.5), Inches(0.24), Inches(8.2), Inches(0.35))
        utf = url_box.text_frame
        up = utf.paragraphs[0]
        up.text = f"URL: {item['url']}"
        up.font.size = Pt(10.5)
        up.font.bold = True
        up.font.color.rgb = CYAN_ACCENT
        up.alignment = PP_ALIGN.RIGHT

        # Slide Main Title
        title_box = s.shapes.add_textbox(Inches(0.6), Inches(0.65), Inches(12.1), Inches(0.6))
        ttf = title_box.text_frame
        tp = ttf.paragraphs[0]
        tp.text = f"{item['num']}. {item['title']}"
        tp.font.size = Pt(18)
        tp.font.bold = True
        tp.font.color.rgb = WHITE

        # Add Screenshot Image (Large, crisp 16:9 view)
        img_file = os.path.join(SCREENSHOTS_DIR, item["img"])
        if os.path.exists(img_file):
            img_left = Inches(0.6)
            img_top = Inches(1.3)
            img_width = Inches(8.4)
            img_height = Inches(4.725) # 16:9 aspect ratio

            # Border frame container behind image
            frame = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, img_left - Inches(0.04), img_top - Inches(0.04), img_width + Inches(0.08), img_height + Inches(0.08))
            frame.fill.solid()
            frame.fill.fore_color.rgb = NAVY_SURFACE
            frame.line.color.rgb = BORDER_COLOR
            frame.line.width = Pt(1.5)

            # Insert Image
            s.shapes.add_picture(img_file, img_left, img_top, img_width, img_height)

        # Right-side Features & Guide Card
        card_left = Inches(9.2)
        card_top = Inches(1.26)
        card_width = Inches(3.533)
        card_height = Inches(4.8)

        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_left, card_top, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = NAVY_SURFACE
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1.5)

        ctb = s.shapes.add_textbox(card_left + Inches(0.2), card_top + Inches(0.15), card_width - Inches(0.4), card_height - Inches(0.3))
        ctf = ctb.text_frame
        ctf.word_wrap = True

        cp0 = ctf.paragraphs[0]
        cp0.text = "SAHIFA IMKONIYATLARI"
        cp0.font.size = Pt(11)
        cp0.font.bold = True
        cp0.font.color.rgb = item["badge_color"]
        cp0.space_after = Pt(8)

        for feat in item["features"]:
            fp = ctf.add_paragraph()
            fp.text = f"• {feat}"
            fp.font.size = Pt(10)
            fp.font.color.rgb = WHITE
            fp.space_after = Pt(6)

        # Bottom Quick Tip Banner
        bot_banner = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.2), Inches(12.133), Inches(0.65))
        bot_banner.fill.solid()
        bot_banner.fill.fore_color.rgb = NAVY_SURFACE
        bot_banner.line.color.rgb = BORDER_COLOR
        bot_banner.line.width = Pt(1)

        bbtf = bot_banner.text_frame
        bbp = bbtf.paragraphs[0]
        bbp.text = f"💡 Boshliq tavsiyasi: Ushbu bo'limda kiritilgan barcha o'zgarishlar real vaqtda bazaga saqlanadi va butun markaz tizimiga ta'sir qiladi."
        bbp.font.size = Pt(10)
        bbp.font.color.rgb = RGBColor(226, 232, 240)
        bbp.alignment = PP_ALIGN.LEFT

    # ==========================================================
    # SLIDE 21: Conclusion & 5 Golden Rules
    # ==========================================================
    s_end = prs.slides.add_slide(blank_layout)
    bg_end = s_end.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg_end.fill.solid()
    bg_end.fill.fore_color.rgb = NAVY_BG
    bg_end.line.fill.background()

    # Header
    e_title_box = s_end.shapes.add_textbox(Inches(1.0), Inches(0.7), Inches(11.3), Inches(1.0))
    e_tf = e_title_box.text_frame
    ep0 = e_tf.paragraphs[0]
    ep0.text = "EDUCO LMS — Boshliq Uchun 5 Oltin Qoida"
    ep0.font.size = Pt(26)
    ep0.font.bold = True
    ep0.font.color.rgb = WHITE

    rules = [
        ("1. Akkauntlarni O'z Vaqtida Yarating", "Admin va o'qituvchilarni 'Foydalanuvchilar' bo'limidan oching. Ularga login uchun telefon raqam va kamida 6 belgili parol bering."),
        ("2. Parollarni Markazlashgan Nazorat Qiling", "'Parollar' bo'limida barcha xodim va talabalar parollarini 1 ta bosishda nusxa olib Telegram orqali yuborishingiz mumkin."),
        ("3. Dars Jadvali To'qnashuvlariga Yo'l Qo'ymang", "'Dars jadvali' bo'limi xona va o'qituvchi vaqtlarini avtomatik tekshiradi, qizil ogohlantirishlarga e'tibor qarating."),
        ("4. Oylik Maoshlarni (Payroll) Oy Yakunida Tasdiqlang", "Har oy oxirida oylik davrini ochib, darsbay yoki fiks hisoblangan maoshlarni tasdiqlab, hisobotni Excelga yuklab oling."),
        ("5. Audit Jurnali Orqali Xavfsizlikni Nazorat Qiling", "Tizimda kim qachon kirdi, nimalarni tahrirladi yoki o'chirdi — 'Audit jurnali'da to'liq saqlanadi.")
    ]

    top_pos = Inches(1.8)
    for title, desc in rules:
        r_box = s_end.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), top_pos, Inches(11.333), Inches(0.85))
        r_box.fill.solid()
        r_box.fill.fore_color.rgb = NAVY_SURFACE
        r_box.line.color.rgb = BLUE_ACCENT
        r_box.line.width = Pt(1)

        r_tf = r_box.text_frame
        r_tf.word_wrap = True
        rp0 = r_tf.paragraphs[0]
        rp0.text = title
        rp0.font.size = Pt(12)
        rp0.font.bold = True
        rp0.font.color.rgb = CYAN_ACCENT
        rp0.space_after = Pt(2)

        rp1 = r_tf.add_paragraph()
        rp1.text = desc
        rp1.font.size = Pt(10)
        rp1.font.color.rgb = WHITE

        top_pos += Inches(0.95)

    # Save PPTX & Export PDF
    out_dir = r"C:\Users\Abdul\Documents\EDUCO_LMS"
    pptx_path = os.path.join(out_dir, "EDUCO_LMS_Boshliq_Qollanma_Skrinshotlar.pptx")
    pdf_path = os.path.join(out_dir, "EDUCO_LMS_Boshliq_Qollanma_Skrinshotlar.pdf")

    prs.save(pptx_path)
    print(f"PPTX muvaffaqiyatli yaratildi: {pptx_path}")

    try:
        import win32com.client
        ppt_app = win32com.client.Dispatch("PowerPoint.Application")
        presentation = ppt_app.Presentations.Open(pptx_path, WithWindow=False)
        presentation.SaveAs(pdf_path, 32) # 32 = ppSaveAsPDF
        presentation.Close()
        ppt_app.Quit()
        print(f"PDF muvaffaqiyatli yaratildi: {pdf_path}")
    except Exception as e:
        print(f"PDF yaratishda xatolik: {e}")

if __name__ == "__main__":
    build_pdf_presentation()
